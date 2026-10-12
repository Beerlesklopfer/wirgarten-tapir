import csv
import datetime
import io
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal

from tapir.utils.exceptions import TapirDataImportException


@dataclass
class TabularData:
    headers: list[str]
    rows: list[dict[str, str]] = field(default_factory=list)


class TabularFileReader:
    """Reads Excel (.xlsx) and CSV files into rows of strings, the format the importers expect.

    Excel cells are typed: dates become ISO dates, whole numbers lose the ".0" Excel adds,
    booleans become "true"/"false". CSV files may use ";" or "," and UTF-8 or Windows-1252.
    """

    EXCEL_EXTENSIONS = (".xlsx", ".xlsm")
    CSV_EXTENSIONS = (".csv", ".txt")

    @classmethod
    def read(cls, file, filename: str) -> TabularData:
        name = filename.lower()
        if name.endswith(cls.EXCEL_EXTENSIONS):
            return cls.read_excel(file)
        if name.endswith(cls.CSV_EXTENSIONS):
            return cls.read_csv(file.read())
        raise TapirDataImportException(
            "Unbekanntes Dateiformat, erlaubt sind Excel (.xlsx) und CSV (.csv)"
        )

    @classmethod
    def read_excel(cls, file) -> TabularData:
        from openpyxl import load_workbook

        try:
            workbook = load_workbook(file, read_only=True, data_only=True)
        except Exception as error:
            raise TapirDataImportException(
                f"Die Excel-Datei kann nicht gelesen werden: {error}"
            )
        try:
            sheet = workbook.worksheets[0]
            rows = sheet.iter_rows(values_only=True)
            raw_headers = next(rows, None)
            if raw_headers is None:
                raise TapirDataImportException("Die Tabelle ist leer")
            headers = [cls.clean_field_name(cls.cell_to_string(h)) for h in raw_headers]
            data = TabularData(headers=headers)
            for values in rows:
                row = {
                    header: cls.cell_to_string(value)
                    for header, value in zip(headers, values)
                    if header
                }
                data.rows.append(row)
            return data
        finally:
            workbook.close()

    @classmethod
    def read_csv(cls, content: bytes) -> TabularData:
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise TapirDataImportException(
                "Die CSV-Datei hat eine unbekannte Kodierung"
            )

        first_line = text.split("\n", 1)[0]
        delimiter = ";" if first_line.count(";") >= first_line.count(",") else ","
        reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
        raw_headers = next(reader, None)
        if raw_headers is None:
            raise TapirDataImportException("Die CSV-Datei ist leer")
        headers = [cls.clean_field_name(h) for h in raw_headers]
        data = TabularData(headers=headers)
        for values in reader:
            data.rows.append(
                {header: value for header, value in zip(headers, values) if header}
            )
        return data

    @staticmethod
    def cell_to_string(value) -> str:
        if value is None:
            return ""
        if isinstance(value, bool):
            return "true" if value else "false"
        if isinstance(value, datetime.datetime):
            if value.time() == datetime.time(0, 0):
                return value.date().isoformat()
            return value.isoformat()
        if isinstance(value, datetime.date):
            return value.isoformat()
        if isinstance(value, float):
            if value.is_integer():
                return str(int(value))
            return format(Decimal(str(value)).normalize(), "f")
        return str(value).strip()

    @staticmethod
    def clean_field_name(name: str | None) -> str:
        if name is None:
            return ""
        # Invisible characters Excel and other tools put into headers: BOM, zero-width spaces,
        # word joiner; non-breaking spaces become regular spaces
        for invisible in ("﻿", "​", "‌", "‍", "⁠"):
            name = name.replace(invisible, "")
        name = name.replace("\xa0", " ")
        return unicodedata.normalize("NFKC", name).strip()
