import io
from dataclasses import dataclass, field
from typing import Callable

from django.db import transaction

from tapir.utils.config import (
    MEMBER_IMPORT_STATUS_CREATED,
    MEMBER_IMPORT_STATUS_SKIPPED,
    MEMBER_IMPORT_STATUS_UPDATED,
)
from tapir.utils.exceptions import DryRunException
from tapir.utils.services.member_importer import MemberImporter
from tapir.utils.services.share_importer import ShareImporter
from tapir.utils.services.subscription_importer import SubscriptionImporter
from tapir.utils.services.tabular_file_reader import TabularData

STATUS_CREATED = "created"
STATUS_UPDATED = "updated"
STATUS_SKIPPED = "skipped"
STATUS_ERROR = "error"

STATUS_FROM_IMPORTER = {
    MEMBER_IMPORT_STATUS_CREATED: STATUS_CREATED,
    MEMBER_IMPORT_STATUS_UPDATED: STATUS_UPDATED,
    MEMBER_IMPORT_STATUS_SKIPPED: STATUS_SKIPPED,
}


@dataclass(frozen=True)
class ImportType:
    key: str
    name: str
    importer: Callable
    required_columns: tuple[str, ...]
    optional_columns: tuple[str, ...]
    example: dict[str, str]
    label_columns: tuple[str, ...]

    @property
    def columns(self) -> tuple[str, ...]:
        return self.required_columns + self.optional_columns


IMPORT_TYPES = {
    import_type.key: import_type
    for import_type in [
        ImportType(
            key="members",
            name="Mitglieder",
            importer=MemberImporter.import_member,
            required_columns=(
                "Nr",
                "Vorname",
                "Nachname",
                "Mailadresse",
                "consent_sepa",
            ),
            optional_columns=(
                "Anrede",
                "Straße",
                "Hausnr.",
                "PLZ",
                "Ort",
                "Telefon",
                "Telefon 2",
                "Geburtstag/Gründungsdatum",
                "IBAN",
                "Kontoinhaber",
                "privacy_consent",
                "Pseudonym",
                "Abholort",
                "AO_gueltig_ab",
                "Zahlungsintervall",
                "Solidarpreis in EUR",
                "Solidarpreis Start_Date",
                "Solidarpreis End_Date",
                "Vereinsmitgliedschafttyp",
                "Vereinsmitgliedschaft Start_Date",
                "Vereinsmitgliedschaft End_Date",
                "Lieblingsbroten",
            ),
            example={
                "Nr": "1",
                "Vorname": "Erika",
                "Nachname": "Musterfrau",
                "Mailadresse": "erika@example.org",
                "consent_sepa": "01.03.2024",
                "Anrede": "Liebe",
                "Straße": "Ackerweg",
                "Hausnr.": "1",
                "PLZ": "12345",
                "Ort": "Musterdorf",
                "Telefon": "+49 170 1234567",
                "Geburtstag/Gründungsdatum": "15.05.1980",
                "IBAN": "DE02120300000000202051",
                "Kontoinhaber": "Erika Musterfrau",
                "privacy_consent": "01.03.2024",
                "Abholort": "Name des Abholorts in Tapir",
                "AO_gueltig_ab": "01.04.2024",
                "Zahlungsintervall": "Monatlich",
            },
            label_columns=("Nr", "Vorname", "Nachname"),
        ),
        ImportType(
            key="shares",
            name="Genossenschaftsanteile",
            importer=ShareImporter.import_shares_single_member,
            required_columns=(
                "Mitgliedsnummer",
                "Bewegungsart (Z,Ü,K)",
                "Datum",
                "Anzahl Anteile",
            ),
            optional_columns=("Übertragungspartner", "Wirkung Kündigung"),
            example={
                "Mitgliedsnummer": "1",
                "Bewegungsart (Z,Ü,K)": "Z",
                "Datum": "01.03.2024",
                "Anzahl Anteile": "2",
            },
            label_columns=("Mitgliedsnummer", "Bewegungsart (Z,Ü,K)", "Anzahl Anteile"),
        ),
        ImportType(
            key="subscriptions",
            name="Verträge",
            importer=SubscriptionImporter.import_subscription,
            required_columns=(
                "Mitgliedernummer",
                "product",
                "Quantity",
                "Vertragsbeginn",
                "Vertragsende",
            ),
            optional_columns=(
                "Email",
                "product_type",
                "Probezeit",
                "Ende Probezeit (Sonntag)",
                "Personalisierter Betrag",
                "consent_vertragsgrundsätze",
                "consent_widerruf",
                "cancellation.ts",
            ),
            example={
                "Mitgliedernummer": "1",
                "product": "M",
                "Quantity": "1",
                "Vertragsbeginn": "01.04.2024",
                "Vertragsende": "31.03.2025",
                "product_type": "Ernteanteile",
                "Probezeit": "nein",
                "consent_vertragsgrundsätze": "01.03.2024",
                "consent_widerruf": "01.03.2024",
            },
            label_columns=("Mitgliedernummer", "product_type", "product"),
        ),
    ]
}


@dataclass
class RowResult:
    line: int
    label: str
    status: str
    message: str = ""


@dataclass
class ImportResult:
    import_type: ImportType
    dry_run: bool
    missing_columns: list[str] = field(default_factory=list)
    unknown_columns: list[str] = field(default_factory=list)
    rows: list[RowResult] = field(default_factory=list)

    def count(self, status: str) -> int:
        return sum(1 for row in self.rows if row.status == status)

    @property
    def counts(self) -> dict[str, int]:
        return {
            status: self.count(status)
            for status in (STATUS_CREATED, STATUS_UPDATED, STATUS_SKIPPED, STATUS_ERROR)
        }

    @property
    def can_be_applied(self) -> bool:
        return not self.missing_columns and self.count(STATUS_ERROR) == 0


class DataImportService:
    @classmethod
    def check_columns(cls, import_type: ImportType, data: TabularData):
        headers = [header for header in data.headers if header]
        missing = [
            column for column in import_type.required_columns if column not in headers
        ]
        unknown = [header for header in headers if header not in import_type.columns]
        return missing, unknown

    @classmethod
    def run(
        cls,
        import_type: ImportType,
        data: TabularData,
        update_existing: bool,
        dry_run: bool,
    ) -> ImportResult:
        """Imports all rows in one transaction, each row in its own savepoint so that a failing
        row is reported and does not stop the others. A dry run rolls everything back and does
        not touch Keycloak."""
        result = ImportResult(import_type=import_type, dry_run=dry_run)
        result.missing_columns, result.unknown_columns = cls.check_columns(
            import_type, data
        )
        if result.missing_columns:
            return result

        keycloak_updates_before = MemberImporter.keycloak_updates_enabled
        MemberImporter.keycloak_updates_enabled = not dry_run
        try:
            with transaction.atomic():
                cls.import_rows(import_type, data, update_existing, result)
                if dry_run or result.count(STATUS_ERROR):
                    raise DryRunException()
        except DryRunException:
            pass
        finally:
            MemberImporter.keycloak_updates_enabled = keycloak_updates_before
        return result

    @classmethod
    def import_rows(
        cls,
        import_type: ImportType,
        data: TabularData,
        update_existing: bool,
        result: ImportResult,
    ):
        for index, row in enumerate(data.rows):
            if not any(str(value).strip() for value in row.values()):
                continue
            line = index + 2  # line 1 holds the headers
            label = " ".join(
                str(row.get(column, "")).strip()
                for column in import_type.label_columns
                if str(row.get(column, "")).strip()
            )
            try:
                with transaction.atomic():
                    status = import_type.importer(
                        row=row, update_existing=update_existing
                    )
                result.rows.append(
                    RowResult(
                        line=line,
                        label=label,
                        status=STATUS_FROM_IMPORTER.get(status, STATUS_SKIPPED),
                    )
                )
            except Exception as error:
                result.rows.append(
                    RowResult(
                        line=line,
                        label=label,
                        status=STATUS_ERROR,
                        message=str(error) or error.__class__.__name__,
                    )
                )

    @classmethod
    def build_template(cls, import_type: ImportType) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = import_type.name
        sheet.append(list(import_type.columns))
        sheet.append(
            [import_type.example.get(column, "") for column in import_type.columns]
        )
        for cell in sheet[1]:
            cell.font = Font(bold=cell.value in import_type.required_columns)
        for index, column in enumerate(import_type.columns, start=1):
            letter = sheet.cell(row=1, column=index).column_letter
            sheet.column_dimensions[letter].width = max(14, len(column) + 2)
            # Text cells: Excel must not turn phone numbers, IBANs or member numbers into numbers
            for row in range(2, 1000):
                sheet.cell(row=row, column=index).number_format = "@"
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()


class ActivationMailService:
    """Keycloak logins and activation mails for imported members (like the command
    send_verify_emails): members without login that have coop shares or a current or future
    subscription."""

    @staticmethod
    def get_members_without_login():
        from tapir.wirgarten.management.commands.send_verify_emails import (
            Command as SendVerifyEmailsCommand,
        )

        return (
            SendVerifyEmailsCommand.get_members_with_at_least_one_share_or_one_relevant_subscription()
            .distinct()
            .order_by("member_no", "last_name")
        )

    @classmethod
    def send(cls, actor) -> tuple[int, list[str]]:
        from tapir.accounts.services.email_verification_service import (
            EmailVerificationService,
        )
        from tapir.accounts.services.keycloak_user_manager import KeycloakUserManager

        cache = {}
        sent = 0
        errors = []
        for member in cls.get_members_without_login():
            try:
                KeycloakUserManager.create_keycloak_user_if_necessary(
                    user=member, initial_password=None, cache=cache
                )
                EmailVerificationService.send_verification_email(
                    user=member, actor=actor, cache=cache
                )
                sent += 1
            except Exception as error:
                errors.append(f"{member}: {error}")
        return sent, errors
